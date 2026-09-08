"""
STEP 1: INPUT & PARSING
------------------------
What it does : Converts the uploaded resume file (PDF/DOCX) into clean plain text,
                and cleans the pasted job description text.
Gets         : Uploaded file object (from Streamlit), raw JD string
Gives        : resume_text (str), jd_text (str)
Uses         : Pure Python (no AI) — this is deterministic file/text handling.
"""

import io
import re
from pypdf import PdfReader
from docx import Document


class ParsingError(Exception):
    """Raised when a resume file cannot be read or produces no usable text."""
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract raw text from a PDF file's bytes."""
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        text_parts.append(page_text)
    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract raw text from a DOCX file's bytes."""
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs]
    # Also pull text from tables, since resumes sometimes use them for layout
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(paragraphs)


def clean_text(raw_text: str) -> str:
    """Normalize whitespace and strip weird characters/line breaks."""
    if not raw_text:
        return ""
    text = raw_text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)          # collapse repeated spaces/tabs
    text = re.sub(r"\n{3,}", "\n\n", text)        # collapse excessive blank lines
    text = "\n".join(line.strip() for line in text.split("\n"))
    return text.strip()


def parse_resume(uploaded_file) -> str:
    """
    Main entry point for Step 1 (resume side).

    Parameters
    ----------
    uploaded_file : a Streamlit UploadedFile object (has .name and .read())

    Returns
    -------
    resume_text : str (cleaned plain text)

    Raises
    ------
    ParsingError if the file type is unsupported or no text could be extracted.
    """
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if filename.endswith(".pdf"):
        raw_text = extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        raw_text = extract_text_from_docx(file_bytes)
    elif filename.endswith(".txt"):
        raw_text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ParsingError(
            f"Unsupported file type: '{filename}'. Please upload a PDF, DOCX, or TXT file."
        )

    cleaned = clean_text(raw_text)

    if not cleaned or len(cleaned) < 30:
        raise ParsingError(
            "Could not extract readable text from this resume. "
            "It may be a scanned image or corrupted file."
        )

    return cleaned


def parse_job_description(jd_raw: str) -> str:
    """
    Main entry point for Step 1 (JD side).

    Parameters
    ----------
    jd_raw : str, raw pasted job description text

    Returns
    -------
    jd_text : str (cleaned plain text)

    Raises
    ------
    ParsingError if the JD is empty or too short to be meaningful.
    """
    cleaned = clean_text(jd_raw)

    if not cleaned or len(cleaned) < 30:
        raise ParsingError(
            "Job description looks empty or too short. Please paste the full JD."
        )

    return cleaned
